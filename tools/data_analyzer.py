"""
数据分析工具
支持读取 Excel/CSV、数据清洗、统计分析、可视化
自动生成分析报告并集成文档生成功能
"""
import os
import json
import pandas as pd
import matplotlib.pyplot as plt
from typing import Optional, Dict, Any, Type
from pydantic import BaseModel, Field
import logging

from config.settings import settings
from .base import BaseTool

logger = logging.getLogger(__name__)

plt.rcParams['font.sans-serif'] = ['SimHei', 'Microsoft YaHei', 'DejaVu Sans']
plt.rcParams['axes.unicode_minus'] = False


class DataAnalyzerInput(BaseModel):#规定参数的校验 5个参数 action是必要参数 剩下的是非必要的
    """数据分析工具输入参数模式"""
    action: str = Field(description="操作类型: read_data, clean_data, analyze, visualize, generate_report")#看你是要用哪个功能
    file_path: Optional[str] = Field(None, description="数据文件路径（Excel/CSV）")
    sheet_name: Optional[str] = Field(None, description="Excel工作表名称")
    session_id: Optional[str] = Field(None, description="会话ID，用于隔离多用户数据")
    auto_generate_doc: Optional[bool] = Field(False, description="是否自动生成Word报告")


class DataAnalyzerTool(BaseTool):#数据分析工具类
    """
    数据分析工具

    功能：
    1. 读取 Excel/CSV 数据
    2. 数据清洗（去重、处理缺失值）
    3. 统计分析（描述性统计、相关性）
    4. 数据可视化（生成图表）
    5. 自动生成 Word 分析报告
    """

    name: str = "data_analyzer"
    description: str = """
    数据分析工具，用于处理 Excel/CSV 文件。
    支持的操作：
    - read_data: 读取数据文件，需要提供 file_path
    - clean_data: 清洗已加载的数据（去重、删除缺失值）
    - analyze: 生成数据统计分析报告
    - visualize: 生成可视化图表
    - generate_report: 生成完整的 Word 分析报告（自动整合分析和可视化结果）
    注意：需要先执行 read_data 才能进行后续操作。
    会话级别的数据隔离：不同 session_id 的数据互不影响。
    """
    args_schema: Type[BaseModel] = DataAnalyzerInput#参数校验等于我们数据分析的入参

    _data_registry: Dict[str, pd.DataFrame] = {}
    _current_file_registry: Dict[str, str] = {}

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self._session_id = None

    @property
    def _data(self) -> Optional[pd.DataFrame]:
        return self._data_registry.get(self._session_id or "default")

    @_data.setter
    def _data(self, value: pd.DataFrame):
        key = self._session_id or "default"
        if value is None:
            self._data_registry.pop(key, None)
        else:
            self._data_registry[key] = value

    @property #@property 的作用：把下面这个方法变成“假装是个变量”。别人写 self._current_file 就能读它，像读普通属性一样，不用写括号。
    def _current_file(self) -> Optional[str]:#先看读取部分（getter） self._session_id or "default"：如果 _session_id 是 None（还没设置会话），就用 "default" 当钥匙。or 在 Python 里的意思是“前面是假值就换后面”，None 算假值，所以会落到 "default"。
        return self._current_file_registry.get(self._session_id or "default")#.get(钥匙)：从储物柜里按钥匙取文件路径。取不到就返回 None，不会报错。

    @_current_file.setter#@_current_file.setter：意思是“给刚才那个属性配一个写入口”。以后代码里写 self._current_file = "xxx" 时，实际执行的是这个方法。
    def _current_file(self, value: str):
        key = self._session_id or "default"#key = ...：同样先算出钥匙（当前用户的）。
        if value is None:#if value is None：如果传进来的是 None，说明“要清掉这个用户记住的文件”，就用 pop(key, None) 把柜子里这条记录删掉。
            self._current_file_registry.pop(key, None)
        else:
            self._current_file_registry[key] = value#否则把新文件名存进柜子：柜子[钥匙] = 文件名。

    def _execute(self, action: str, file_path: Optional[str] = None,
                    sheet_name: Optional[str] = None, session_id: Optional[str] = None,
                 auto_generate_doc: bool = False, **kwargs) -> str:#路由操作
        """
        执行数据分析操作
        """
        self._session_id = session_id or "default"

        if action == "read_data":
            return self._read_data(file_path, sheet_name)
        elif action == "clean_data":
            return self._clean_data()
        elif action == "analyze":
            return self._analyze()
        elif action == "visualize":
            return self._visualize()
        elif action == "generate_report":
            return self._generate_full_report(auto_generate_doc)
        else:
            return f"未知操作: {action}，支持的操作: read_data, clean_data, analyze, visualize, generate_report"

    def _get_analysis_summary(self) -> str:#获取分析结果的摘要
        """获取分析结果的摘要，用于填充报告"""
        if self._data is None:
            return "无数据"

        df = self._data
        summary_parts = []

        summary_parts.append(f"数据规模: {len(df)} 行, {len(df.columns)} 列")
        summary_parts.append(f"列名: {', '.join(df.columns.astype(str))}")

        numeric_cols = df.select_dtypes(include=['number']).columns
        if len(numeric_cols) > 0:
            summary_parts.append(f"\n数值列 ({len(numeric_cols)} 个): {', '.join(numeric_cols)}")
            for col in numeric_cols:
                summary_parts.append(f"  - {col}: 均值={df[col].mean():.2f}, 最大={df[col].max():.2f}, 最小={df[col].min():.2f}")

        text_cols = df.select_dtypes(include=['object']).columns
        if len(text_cols) > 0:
            summary_parts.append(f"\n文本列 ({len(text_cols)} 个): {', '.join(text_cols)}")

        return "\n".join(summary_parts)

    def _generate_full_report(self, auto_generate_doc: bool = False) -> str:#Word报告自动生成
        """生成完整分析报告"""
        if self._data is None:
            return "请先使用 read_data 读取数据文件"

        try:
            results = []#拼文字报告 还是建立空列表这个老方法
            results.append("=" * 50)
            results.append("数据分柝报告")
            results.append("=" * 50)

            df = self._data
            results.append(f"\n【数据概览】")
            results.append(f"总行数: {len(df)}")
            results.append(f"总列数: {len(df.columns)}")

            numeric_cols = df.select_dtypes(include=['number']).columns
            if len(numeric_cols) > 0:
                results.append(f"\n【数值分析】")
                stats = df[numeric_cols].describe()
                results.append(stats.to_string())

                if len(numeric_cols) > 1:
                    results.append(f"\n【相关性分析】")
                    corr = df[numeric_cols].corr().round(2)
                    results.append(corr.to_string())

            text_cols = df.select_dtypes(include=['object']).columns
            if len(text_cols) > 0:
                results.append(f"\n【分类统计】")
                for col in text_cols[:5]:
                    value_counts = df[col].value_counts().head(5)
                    results.append(f"{col}: {dict(value_counts)}")

            report_text = "\n".join(results)

            if auto_generate_doc:#生成 Word
                try:
                    from docx import Document
                    from docx.shared import Pt, RGBColor
                    from docx.enum.text import WD_ALIGN_PARAGRAPH
                    from datetime import datetime

                    doc = Document()
                    doc.add_heading('数据分析报告', level=0)

                    doc.add_paragraph(f'生成时间: {datetime.now().strftime("%Y年%m月%d日 %H:%M")}')
                    doc.add_paragraph(f'数据文件: {self._current_file or "未知"}')
                    doc.add_paragraph(f'数据规模: {len(df)} 行 x {len(df.columns)} 列')

                    doc.add_heading('一、数据概览', level=1)
                    doc.add_paragraph(f'共 {len(df)} 条记录，{len(df.columns)} 个字段。')
                    doc.add_paragraph(f'字段列表: {", ".join(df.columns.astype(str))}')

                    if len(numeric_cols) > 0:
                        doc.add_heading('二、数值分析', level=1)
                        doc.add_paragraph('描述性统计:')
                        doc.add_paragraph(stats.to_string())

                        if len(numeric_cols) > 1:
                            doc.add_heading('三、相关性分析', level=1)
                            doc.add_paragraph('相关系数矩阵:')
                            doc.add_paragraph(corr.to_string())

                    doc.add_heading('四、分类统计', level=1)
                    for col in text_cols[:5]:
                        doc.add_paragraph(f'{col}分布:')
                        vc = df[col].value_counts().head(5)
                        for val, count in vc.items():
                            doc.add_paragraph(f'  {val}: {count}', style='List Bullet')

                    viz_dir = settings.VISUALIZATION_OUTPUT_DIR
                    heatmap_path = os.path.join(viz_dir, "correlation_heatmap.png")
                    if os.path.exists(heatmap_path):
                        doc.add_heading('五、可视化图表', level=1)
                        doc.add_picture(heatmap_path, width=6*1024*1024/1024)
                    #保存
                    output_dir = settings.DOCUMENT_OUTPUT_DIR
                    os.makedirs(output_dir, exist_ok=True)
                    file_name = f"数据分析报告_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
                    file_path = os.path.join(output_dir, file_name)
                    doc.save(file_path)
                    report_text += f"\n\n✅ Word报告已生成: {file_path}"
                except ImportError:
                    report_text += "\n\n⚠️ 报告生成依赖 python-docx，请确保已安装"
                except Exception as e:
                    report_text += f"\n\n⚠️ Word报告生成失败: {str(e)}"

            return report_text

        except Exception as e:
            logger.error(f"生成报告失败: {e}")
            return f"生成报告失败: {str(e)}"
    
    def _read_data(self, file_path: Optional[str], sheet_name: Optional[str]) -> str:#读取数据 参数有两个 一个是文件的路径 Optional[str] 表示“可以是字符串，也可以是 None（没传）” 一个是Excel 工作表名，同样可传可不传 这个方法返回的是一个str字符串
        """读取数据文件"""#它返回的是字符串（比如“成功读取数据文件…”“文件不存在”），而不是把数据表格本身返回出去。因为数据已经通过 self._data = ... 存进“会话储物柜”了，后面的清洗、分析、画图方法都要从柜子里取，返回值只是给用户/Agent 看的提示文字。
        if not file_path:
            file_path = settings.DEFAULT_DATA_PATH#如果用户没有传路径 就使用默认的路径
        
        # 如果指定路径不存在，尝试在上传目录中查找
        original_path = file_path#先把原始路径存一份，因为后面可能被改写，报错时要用原路径告诉用户“你给的那个文件不存在”。
        if not os.path.exists(file_path):#检查这个路径（file_path）的文件存不存在。
            upload_dir = "data/uploads"
            file_name = os.path.basename(file_path)#从完整路径里取出文件名部分，结果是 sales.xlsx。
            upload_path = os.path.join(upload_dir, file_name)#把目录和文件名拼成完整路径，结果是 data/uploads/sales.xlsx。用 join 而不是手动加 /，是因为它能自动处理分隔符（Windows 是 \，Mac/Linux 是 /）。
            if os.path.exists(upload_path):
                file_path = upload_path
            else:
                return f"文件不存在: {original_path}"
        
        try:
            ext = os.path.splitext(file_path)[1].lower()#获取上传文件的后缀 ext
            if ext in ['.xlsx', '.xls']:#如果文件后缀里面存在xlsx 或 xls
                data = pd.read_excel(file_path, sheet_name=sheet_name)#用pd去读取excel
                # 如果返回的是字典（多个工作表），选择第一个工作表
                if isinstance(data, dict) and data:
                    self._data = list(data.values())[0]
                else:
                    self._data = data
            elif ext == '.csv':#如果后缀是csv 那就用pd去读取csv文件
                self._data = pd.read_csv(file_path)
            else:
                return f"不支持的文件格式: {ext}，请使用 Excel 或 CSV 文件"#如果上传的两个都不是 就返回不支持文件类型
            
            # 确保_data是DataFrame对象
            if not isinstance(self._data, pd.DataFrame):#pd.DataFrame 就是 pandas 里的“表格对象”，也就是一行行一列列的数据。这里相当于质检：读进来的必须是一张正常表格，否则报错。
                return f"读取数据失败: 数据格式不正确，预期DataFrame对象"
            
            self._current_file = file_path#把最终成功读取的路径存进柜子（用我们上节讲的 property），后面生成报告时能知道数据来自哪个文件。
            rows, cols = self._data.shape#返回 (行数, 列数) 元组。rows, cols = ... 是元组解包，一次性把行数、列数分别赋给两个变量。
            return f"成功读取数据文件: {file_path}\n数据规模: {rows} 行, {cols} 列\n列名: {', '.join(self._data.columns.astype(str))}"
        except Exception as e:
            logger.error(f"读取数据失败: {e}")
            return f"读取数据失败: {str(e)}"
    
    def _clean_data(self) -> str:#清洗数据
        """清洗数据"""
        if self._data is None:#如果没有上传文件 柜子里面没有数据 返回要求
            return "请先使用 read_data 读取数据文件"
        
        try:
            original_rows = len(self._data)#先把原始行数记下来，是为了最后能告诉用户“我帮你清理掉了多少行”。
            
            # 删除重复行
            self._data = self._data.drop_duplicates()
            
            # 删除全为空的行
            self._data = self._data.dropna(how='all')
            
            # 对于数值列，用中位数填充缺失值
            numeric_cols = self._data.select_dtypes(include=['number']).columns
            for col in numeric_cols:
                # 跳过空列或全为 NaN 的列
                if self._data[col].notna().any():
                    median_val = self._data[col].median()
                    self._data[col] = self._data[col].fillna(median_val)
            
            # 对于文本列，用"未知"填充
            text_cols = self._data.select_dtypes(include=['object']).columns
            for col in text_cols:
                self._data[col] = self._data[col].fillna("未知")
            #算结果并汇报
            new_rows = len(self._data)
            removed = original_rows - new_rows
            return f"数据清洗完成。原始行数: {original_rows}，清洗后行数: {new_rows}，移除 {removed} 行无效数据。"
        except Exception as e:
            logger.error(f"数据清洗失败: {e}")
            return f"数据清洗失败: {str(e)}"
    
    def _analyze(self) -> str:#统计分析与可视化方法
        """统计分析"""
        if self._data is None:
            return "请先使用 read_data 读取数据文件"#还是先分析有没有文件 没有读取文件先去读取文件
        
        try:
            df = self._data#给柜子中的数据起一个别名 叫做df
            result_parts = []#先建一个空列表，把要输出的每一行 append 进去，最后用 "\n".join(...) 拼成一大段文字。 好处是每一行独立管理，不会像疯狂用 += 拼字符串那样越拼越乱。你以后写“要输出多段内容”的功能，直接抄这个模式。
            
            # 基本统计信息
            result_parts.append("=== 数据概览 ===")
            result_parts.append(f"总行数: {len(df)}")
            result_parts.append(f"总列数: {len(df.columns)}")
            result_parts.append(f"列名: {', '.join(df.columns.astype(str))}")
            
            # 数值列统计
            numeric_cols = df.select_dtypes(include=['number']).columns#挑出数字列（比如 单价、销售数量、折扣率、销售额）
            if len(numeric_cols) > 0:
                result_parts.append("\n=== 数值列统计 ===")
                stats = df[numeric_cols].describe().to_string()#只取这些数字列，组成一张子表。
                result_parts.append(stats)
                
                # 相关性分析（如果有多列） 相关性矩阵
                if len(numeric_cols) > 1:
                    result_parts.append("\n=== 相关性矩阵 ===")
                    corr = df[numeric_cols].corr().round(2).to_string()
                    result_parts.append(corr)
            
            # 文本列统计
            text_cols = df.select_dtypes(include=['object']).columns
            if len(text_cols) > 0:
                result_parts.append("\n=== 文本列统计 ===")
                for col in text_cols:
                    unique_count = df[col].nunique()
                    result_parts.append(f"{col}: 唯一值数量 {unique_count}")
            
            return "\n".join(result_parts)
        except Exception as e:
            logger.error(f"数据分析失败: {e}")
            return f"数据分析失败: {str(e)}"
    
    def _visualize(self) -> str:#生成可视化图表方法
        """生成可视化图表"""
        if self._data is None:
            return "请先使用 read_data 读取数据文件"#依旧先查数据存不存在
        
        try:
            df = self._data
            numeric_cols = df.select_dtypes(include=['number']).columns
            
            if len(numeric_cols) == 0:#如果数据中没有列则无法生成可视化的图表
                return "数据中没有数值列，无法生成可视化图表"
            
            output_dir = settings.VISUALIZATION_OUTPUT_DIR#输出文件的存储
            os.makedirs(output_dir, exist_ok=True)
            
            generated_files = []#创建收集文件路径的列表
            
            # 为每个数值列生成直方图
            for col in numeric_cols[:3]:  # 限制最多3个图表
                plt.figure(figsize=(10, 6))#新建一张空白画布，宽 10、高 6（英寸）。每次循环都新建，否则会画到上一张图上。
                plt.hist(df[col].dropna(), bins=20, edgecolor='black', alpha=0.7)
                plt.title(f'{col} 分布直方图')
                plt.xlabel(col)
                plt.ylabel('频数')
                plt.grid(True, alpha=0.3)
                
                file_name = f"{col}_histogram.png"
                file_path = os.path.join(output_dir, file_name)
                plt.savefig(file_path, dpi=150, bbox_inches='tight')
                plt.close()
                generated_files.append(file_path)
            
            # 如果有多个数值列，生成相关性热力图
            if len(numeric_cols) > 1:
                plt.figure(figsize=(10, 8))#plt.figure() / plt.savefig() / plt.close() 三件套：开画布 → 存文件 → 关画布，每次画图都要完整走一遍。
                corr = df[numeric_cols].corr()
                plt.imshow(corr, cmap='coolwarm', aspect='auto')
                plt.colorbar(label='相关系数')
                plt.xticks(range(len(corr.columns)), corr.columns, rotation=45, ha='right')
                plt.yticks(range(len(corr.columns)), corr.columns)
                plt.title('特征相关性热力图')
                
                file_path = os.path.join(output_dir, "correlation_heatmap.png")
                plt.savefig(file_path, dpi=150, bbox_inches='tight')
                plt.close()
                generated_files.append(file_path)
            
            return f"可视化图表生成成功，共生成 {len(generated_files)} 个图表:\n" + "\n".join(generated_files)
        except Exception as e:
            logger.error(f"可视化生成失败: {e}")
            return f"可视化生成失败: {str(e)}"